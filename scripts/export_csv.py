"""The Excel workbook sent with the report emails (one .xlsx, one sheet per table).

  python scripts/export_csv.py weekly  <out_dir>
  python scripts/export_csv.py monthly <out_dir> [YYYY-MM]
  python scripts/export_csv.py quarterly <out_dir> [YYYY-Qn]

Writes vessel_data_<kind>_<period>.xlsx and prints its path. Sheets: Ship summary,
Positions (period), Port calls, Full history.
"""
import csv
import os
import sys
from datetime import timedelta, timezone

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import haversine_nm, load_roster, now_utc, place_label, places_visited, read_positions  # noqa: E402
from weekly_report import period_progress, load_fixes, period_bounds  # noqa: E402
from insights import port_calls  # noqa: E402

SGT = timezone(timedelta(hours=8))


def sgt(dt):
    return dt.astimezone(SGT).strftime("%Y-%m-%d %H:%M") if dt else ""


def write(path, header, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    return path


def write_xlsx(path, sheets):
    """One Excel workbook: sheets = [(name, header, rows)]. Bold frozen header, filters, sensible widths."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(name[:31])
        ws.append(header)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0B4F7A")
            c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), start=1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 60)
    wb.save(path)
    return path


def position_rows(fixes_by_id, roster, start=None, end=None):
    by_id = {v["id"]: v for v in roster}
    rows = []
    for vid, fixes in fixes_by_id.items():
        v = by_id.get(vid, {})
        for f in fixes:
            if (start and f["t"] < start) or (end and f["t"] > end):
                continue
            rows.append((f["t"], [
                sgt(f["t"]), vid, v.get("vessel_name", ""), v.get("category", ""),
                place_label(f["lat"], f["lon"]), f["lat"], f["lon"],
                "" if f["speed_kn"] is None else f["speed_kn"], f["nav_status"], f["destination"],
            ]))
    rows.sort(key=lambda r: (r[0], r[1][1]))
    return [r[1] for r in rows]


POS_HEADER = ["Time (SGT)", "ID", "Vessel", "Category", "Place", "Latitude", "Longitude",
              "Speed (kn)", "Status", "Destination (AIS)"]


def main():
    kind, out_dir = sys.argv[1], sys.argv[2]
    month = sys.argv[3] if len(sys.argv) > 3 and sys.argv[3] else None
    os.makedirs(out_dir, exist_ok=True)
    ref = now_utc()
    start, end, _period, tag, ndays = period_bounds(kind, ref, month)
    as_of = min(end, ref)
    roster = load_roster()
    fixes = load_fixes(read_positions())
    label = f"{kind}_{tag}"

    sheets = [("Positions (period)", POS_HEADER, position_rows(fixes, roster, start, as_of))]

    summary = []
    for v in roster:
        hist = [f for f in fixes.get(v["id"], []) if f["t"] <= as_of]
        period = [f for f in hist if f["t"] >= start]
        last = hist[-1] if hist else None
        dist = sum(haversine_nm(a["lat"], a["lon"], b["lat"], b["lon"]) for a, b in zip(period, period[1:]))
        prog, _ = period_progress(v, as_of.astimezone(SGT).date())
        summary.append([
            v["id"], v["vessel_name"], v.get("company", ""), v.get("category", ""), v["imo"], v["mmsi"],
            v.get("start_date", ""), v.get("end_date", ""), prog,
            sgt(last["t"]) if last else "", place_label(last["lat"], last["lon"]) if last else "",
            last["lat"] if last else "", last["lon"] if last else "",
            f"{len({f['t'].astimezone(SGT).date() for f in period})}/{ndays}", len(period), round(dist),
            " -> ".join(places_visited(period)),
        ])
    sheets.insert(0, ("Ship summary", [
        "ID", "Vessel", "Company", "Category", "IMO", "MMSI", "First sign-on", "Expected finish (182 days on board)",
        "Progress (days on board)", "Last seen (SGT)", "Last place", "Last latitude", "Last longitude",
        "Days heard", "Positions in period", "Distance logged (nm)", "Places visited (within 30 nm)",
    ], summary))

    calls = []
    for v in roster:
        hist = [f for f in fixes.get(v["id"], []) if f["t"] <= as_of]
        for c in port_calls(hist):
            if c["end"] >= start and c["start"] <= as_of:
                calls.append([v["id"], v["vessel_name"], c["kind"], c["place"], c["country"],
                              sgt(c["start"]), sgt(c["end"]), sgt(c["departed"]), c["hours"],
                              round((c["departed"] - c["start"]).total_seconds() / 3600, 1) if c["departed"] else "",
                              c["lat"], c["lon"]])
    sheets.append(("Port calls", [
        "ID", "Vessel", "Type", "Place", "Country", "First logged stopped (SGT)", "Last logged stopped (SGT)",
        "Logged under way again (SGT)", "Stopped at least (h)", "Stopped at most (h)", "Latitude", "Longitude",
    ], calls))

    sheets.append((f"Full history to {ref.astimezone(SGT):%d %b %Y}", POS_HEADER, position_rows(fixes, roster)))
    print(write_xlsx(os.path.join(out_dir, f"vessel_data_{label}.xlsx"), sheets))


if __name__ == "__main__":
    main()
